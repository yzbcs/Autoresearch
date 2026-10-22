import json
import os
from pathlib import Path
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
import unittest
from unittest.mock import patch

import pymupdf
from PIL import Image
from harness.providers import OpenAICompatible, Gemini
from harness.standalone import Session, prepare, loop, TOOLS, verify_frozen


class StandaloneTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name); self.inputs = self.base / 'inputs'; self.inputs.mkdir()
        with pymupdf.open() as d:
            d.new_page().insert_text((40, 50), 'Synthetic transport fixture, not a scientific validation')
            d.save(self.inputs / 'paper-without-main-figure.pdf')
        for i in (1, 2, 3):
            Image.new('RGB', (20, 20), (i*60, 40, 90)).save(self.inputs / f'style-reference-{i:02d}.png')
        self.workspace = prepare(self.base / 'workspace', self.inputs, 'fixture-only')

    def test_workspace_escape_and_symlink_rejected(self):
        s = Session(self.workspace)
        with self.assertRaises(ValueError): s.call('read_text', {'path': '../outside'})
        (self.workspace / 'work/escape').symlink_to(self.base, target_is_directory=True)
        with self.assertRaises(ValueError): s.call('write_file', {'path': 'work/escape/file', 'content': 'wrong'})

    def test_inputs_and_guides_cannot_be_overwritten_by_model_tool(self):
        with self.assertRaises(ValueError):
            Session(self.workspace).call('write_file', {'path': 'run/inputs/paper-text.txt', 'content': 'tamper'})

    def test_batch_tool_results_precede_all_pixel_user_messages(self):
        image = str(self.inputs / 'style-reference-01.png')
        msgs = [{'role': 'assistant', 'text': '', 'calls': [
                    {'id': 'a', 'name': 'view_image', 'arguments': {'path': 'a'}},
                    {'id': 'b', 'name': 'view_image', 'arguments': {'path': 'b'}}]},
                {'role': 'tool', 'call_id': 'a', 'name': 'view_image', 'result': {}, 'image_path': image},
                {'role': 'tool', 'call_id': 'b', 'name': 'view_image', 'result': {}, 'image_path': image}]
        wire = OpenAICompatible({'model': 'fixture'}).payload(msgs, TOOLS)['messages']
        self.assertEqual([m['role'] for m in wire], ['assistant', 'tool', 'tool', 'user'])
        self.assertEqual(sum(p['type'] == 'image_url' for p in wire[-1]['content']), 2)

    def test_gemini_thought_signature_and_image_survive_tool_roundtrip(self):
        parts = [{'functionCall': {'name': 'view_image', 'args': {'path': 'a'}}, 'thoughtSignature': 'test-signature'}]
        wire = Gemini({'model': 'fixture'}).payload([
            {'role': 'system', 'text': 'policy'},
            {'role': 'assistant', 'provider_parts': parts},
            {'role': 'tool', 'name': 'view_image', 'result': {'sha256': 'test'},
             'image_path': str(self.inputs / 'style-reference-01.png')}], TOOLS)
        self.assertEqual(wire['contents'][0]['parts'], parts)
        self.assertIn('inlineData', wire['contents'][1]['parts'][1])
        self.assertEqual(wire['systemInstruction']['parts'][0]['text'], 'policy')
        decl = wire['tools'][0]['functionDeclarations']
        self.assertNotIn('parameters', decl[0])
        self.assertIn('parametersJsonSchema', decl[1])

    def test_missing_key_fails_without_request(self):
        with patch.dict(os.environ, {}, clear=True), self.assertRaisesRegex(RuntimeError, 'empty'):
            OpenAICompatible({'model': 'x', 'endpoint': 'http://unused', 'api_key_env': 'NONEXISTENT'}).complete([], TOOLS)

    def test_compatible_thinking_state_roundtrips_after_tool_calls(self):
        messages = [{'role': 'assistant', 'text': 'a', 'provider_reasoning': 'provider-test-state'},
                    {'role': 'user', 'text': 'continue'}]
        wire = OpenAICompatible({'model': 'fixture'}).payload(messages, TOOLS)['messages']
        self.assertEqual(wire[0]['reasoning_content'], 'provider-test-state')

    def test_gemini_real_http_vision_function_response_roundtrip(self):
        bodies=[]; headers=[]
        parts=[{'functionCall': {'id': 'actual-call-id', 'name': 'view_image', 'args': {'path': 'input.png'}},
                'thoughtSignature': 'opaque-fixture-signature'}]
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                bodies.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
                headers.append(self.headers.get('x-goog-api-key'))
                self.send_response(200); self.send_header('Content-Type','application/json');self.end_headers()
                self.wfile.write(json.dumps({'candidates':[{'content':{'parts':parts}}],
                                            'usageMetadata':{'totalTokenCount':0}}).encode())
            def log_message(self,*args):pass
        server=HTTPServer(('127.0.0.1',0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            provider=Gemini({'model':'fixture','endpoint':f'http://127.0.0.1:{server.server_port}',
                             'api_key_env':'FIXTURE_KEY'})
            with patch.dict(os.environ,{'FIXTURE_KEY':'local-fixture-not-real-credential'}):
                messages=[{'role':'system','text':'fixture'},{'role':'user','text':'view'}]
                response=provider.complete(messages,TOOLS)
                messages.append({'role':'assistant',**response})
                messages.append({'role':'tool','name':'view_image','result':{'fixture':True},
                                 'provider_call_id':response['calls'][0]['provider_call_id'],
                                 'image_path':str(self.inputs/'style-reference-01.png')})
                provider.complete(messages,TOOLS)
        finally:
            server.shutdown();server.server_close();thread.join()
        self.assertEqual(headers,['local-fixture-not-real-credential']*2)
        self.assertEqual(bodies[1]['contents'][-2]['parts'],parts)
        result=bodies[1]['contents'][-1]['parts']
        self.assertEqual(result[0]['functionResponse']['id'],'actual-call-id')
        self.assertTrue(result[1]['inlineData']['data'])

    def test_real_http_loop_to_actual_pptx_render_and_freeze(self):
        """Local scripted endpoint exercises transport, never claims model quality."""
        if not os.environ.get('DRAWING_SOFFICE'):
            self.skipTest('Provide public LibreOffice to run real renderer integration')
        builder = '''from pptx import Presentation
from pptx.util import Inches
from pathlib import Path
import sys
p=Presentation(); p.slide_width=Inches(16); p.slide_height=Inches(9)
s=p.slides.add_slide(p.slide_layouts[6])
s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text='Transport fixture only'
p.save(Path(sys.argv[1])/'candidate.pptx')
'''
        steps = [('read_text', {'path': 'MODEL.txt'})]
        steps += [('view_image', {'path': f'run/inputs/style-reference-{i:02d}.png'}) for i in (1,2,3)]
        steps += [('pdf_page', {'page': 1})]
        for name, content in [('understanding', 'fixture not science'), ('representation', 'fixture text'),
                              ('decision', 'fixture decision'), ('builder', builder), ('caption', 'fixture caption')]:
            steps.append(('write_file', {'path': f'work/{name}.'+('py' if name == 'builder' else 'md'), 'content': content}))
        steps += [('execute', {k: 'work/'+v for k,v in {
            'builder': 'builder.py', 'decision': 'decision.md', 'understanding': 'understanding.md', 'representation': 'representation.md'}.items()}),
            ('view_image', {'path': 'run/iterations/000/render/slide-01.png'}),
            ('write_file', {'path': 'work/review.md', 'content': 'Scripted fixture review, not model visual acceptance'}),
            ('review', {'iteration': '000', 'path': 'work/review.md'}),
            ('finish', {'iteration': '000', 'caption': 'work/caption.md'})]
        bodies = []
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                bodies.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
                name, a = steps[len(bodies)-1]
                payload = {'model': 'local-scripted-fixture', 'usage': {'total_tokens': 0}, 'choices': [{'message': {
                    'content': '', 'tool_calls': [{'id': f'c{len(bodies)}', 'type': 'function',
                    'function': {'name': name, 'arguments': json.dumps(a)}}]}}]}
                self.send_response(200); self.send_header('Content-Type', 'application/json'); self.end_headers()
                self.wfile.write(json.dumps(payload).encode())
            def log_message(self, *args): pass
        server = HTTPServer(('127.0.0.1', 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            loop(self.workspace, {'provider': 'openai-compatible', 'endpoint': f'http://127.0.0.1:{server.server_port}',
                                 'model': 'local-scripted-fixture', 'require_key': False}, 30)
        finally:
            server.shutdown(); server.server_close(); thread.join()
        self.assertEqual(len(bodies), len(steps))
        self.assertTrue(any('data:image/png;base64,' in json.dumps(b) for b in bodies))
        self.assertTrue((self.workspace / 'delivery/main-figure.pptx').exists())
        self.assertTrue((self.workspace / 'delivery/main-figure.pdf').exists())
        manifest = json.loads((self.workspace / 'freeze.json').read_text())['files']
        from harness.model_led_drawing import digest
        self.assertTrue(all(digest(self.workspace / n) == sha for n,sha in manifest.items()))
        with self.assertRaisesRegex(ValueError, 'frozen'):
            Session(self.workspace).call('write_file', {'path': 'work/mutated.md', 'content': 'no'})
        (self.workspace / 'new-history-file.txt').write_text('unrecorded new material')
        self.assertEqual(verify_frozen(self.workspace)['unexpected_files'], ['new-history-file.txt'])


if __name__ == '__main__': unittest.main()
