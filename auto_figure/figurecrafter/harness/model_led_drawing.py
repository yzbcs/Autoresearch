"""File/launch/render assistance for an active model, with an external outer loop.

No paper interpretation, style selection, layout generation or automatic repair.
Run with the project Python: python -m harness.model_led_drawing --help.
"""
import argparse
import hashlib
import json
import os
import re
from pathlib import Path
import shutil
import subprocess
import tempfile
import time
from datetime import datetime, timezone

import pymupdf
from pptx import Presentation
from scripts.pptx_object_evidence import extract as object_evidence


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write(path, value):
    path = Path(path)
    with path.open('x') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)


def inventory(path):
    """Expose observed native objects, without converting them to a style policy."""
    prs = Presentation(path)
    def describe(shape):
        item = {'name': shape.name, 'type': str(shape.shape_type),
                'bounds_mm': [round(v / 36000, 3) for v in
                              (shape.left, shape.top, shape.width, shape.height)]}
        item['native_colors'] = [dict(node.attrib) for node in
                                shape._element.xpath('.//a:srgbClr')]
        if shape.has_text_frame:
            item['text'] = shape.text
            item['runs'] = [{'text': r.text, 'font': r.font.name,
                             'size_pt': r.font.size.pt if r.font.size else None,
                             'bold': r.font.bold, 'italic': r.font.italic}
                            for p in shape.text_frame.paragraphs for r in p.runs]
        if hasattr(shape, 'shapes'):
            item['children'] = [describe(s) for s in shape.shapes]
        return item
    return {'canvas_mm': [prs.slide_width / 36000, prs.slide_height / 36000],
            'slides': [[describe(s) for s in slide.shapes] for slide in prs.slides],
            'interpretation': 'Model must inspect native PPTX and rendered pixels.'}


def render(pptx, out):
    out = Path(out)
    out.mkdir()
    binary = os.environ.get('DRAWING_SOFFICE') or shutil.which('soffice') or shutil.which('libreoffice')
    if not binary:
        raise RuntimeError('LibreOffice is required for actual PPTX rendering')
    with tempfile.TemporaryDirectory(prefix='drawing-lo-') as profile:
        argv = [binary, '-env:UserInstallation=' + Path(profile).as_uri(),
                '--headless', '--convert-to', 'pdf', '--outdir', str(out.resolve()),
                str(Path(pptx).resolve())]
        result = subprocess.run(argv, capture_output=True, text=True, timeout=120)
    write(out / 'conversion.json', {'argv': argv, 'exit_code': result.returncode,
                                   'stdout': result.stdout, 'stderr': result.stderr})
    pdf = out / (Path(pptx).stem + '.pdf')
    if result.returncode or not pdf.exists():
        raise RuntimeError('Actual PPTX render failed; retained conversion evidence')
    with pymupdf.open(pdf) as doc:
        for i, page in enumerate(doc):
            page.get_pixmap(matrix=pymupdf.Matrix(1600 / page.rect.width,
                                                   1600 / page.rect.width)).save(
                out / f'slide-{i + 1:02d}.png')
            page.get_pixmap(matrix=pymupdf.Matrix(1063 / page.rect.width,
                                                   1063 / page.rect.width)).save(
                out / f'slide-{i + 1:02d}-publication.png')
    return {p.name: digest(p) for p in out.iterdir() if p.is_file()}


def verify(root):
    manifest = json.loads((root / 'inputs.json').read_text())
    for name, sha in manifest['hashes'].items():
        if digest(root / 'inputs' / name) != sha:
            raise ValueError('Input snapshot changed: ' + name)
    evidence = root / 'input-evidence.json'
    if evidence.exists():
        for name, sha in json.loads(evidence.read_text())['hashes'].items():
            if digest(root / name) != sha:
                raise ValueError('Input evidence changed: ' + name)
    return manifest


IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp'}
DEFAULT_REQUEST = ('Generate the missing-method main figure from the supplied PDF and '
                   'style references. One editable 16:9 PPTX slide with PDF/PNG export. '
                   'The active model owns science, composition, drawing and visual review. '
                   'No global presentation title. Quality is judged by the model, not file checks.')


def figure_contract(pptx, contract):
    """Check dimensions/page count only; never resize or infer figure quality."""
    prs = Presentation(pptx)
    expected_count = contract.get('slide_count', 1)
    if len(prs.slides) != expected_count:
        raise ValueError(f'Expected {expected_count} main-figure slide')
    width, height = prs.slide_width, prs.slide_height
    aspect = contract.get('aspect_ratio', [16, 9])
    if width <= 0 or height <= 0 or abs(width * aspect[1] - height * aspect[0]) > max(aspect):
        raise ValueError('Main figure must be 16:9; model must recompose, not auto-stretch')
    return {'slide_count': len(prs.slides), 'width_emu': width,
            'height_emu': height, 'aspect_ratio': aspect,
            'scope': 'geometry only, no scientific or aesthetic acceptance'}


def prepare(root, pdf, reference, target_id=None, reference_id=None, request=None,
            reference_native=None, mode='production', guide_directory=None,
            outer_loop='human'):
    """Production takes reference images only; legacy native inspection is explicit."""
    root, pdf = Path(root), Path(pdf)
    if outer_loop not in {'human', 'model'}:
        raise ValueError('Outer loop must be human or model')
    guide_base = Path(guide_directory) if guide_directory else Path(__file__).resolve().parents[1] / 'guides'
    guide_names = ['MODEL_LED_DRAWING.zh-CN.md', 'OBJECT_COMPOSITION.zh-CN.md',
                   'STYLE_LEARNING.zh-CN.md', 'REFERENCE_JOINT_LEARNING.zh-CN.md']
    for name in guide_names:
        if not (guide_base / name).is_file() or not (guide_base / name).read_text().strip():
            raise ValueError('Missing or empty drawing guide: ' + name)
    references = reference if isinstance(reference, (list, tuple)) else [reference]
    references = [Path(v) for v in references if v is not None]
    if not references:
        raise ValueError('At least one style reference is required')
    identities = ([None] * len(references) if reference_id is None else
                  reference_id if isinstance(reference_id, (list, tuple)) else [reference_id])
    if len(references) != len(identities):
        raise ValueError('Every supplied reference must have its corresponding paper identity')
    if target_id and any(identity and identity.strip() == target_id.strip()
                         for identity in identities):
        raise ValueError('Style reference must be from a different paper')
    # Several figures from the same other paper are valid distinct visual references.
    companions = dict(reference_native or {})
    if any(not isinstance(i, int) or i < 1 or i > len(references) for i in companions):
        raise ValueError('Native companion index must identify an existing reference')
    if mode not in {'production', 'legacy-reference-inspection'}:
        raise ValueError('Unknown input mode')
    if mode == 'production' and (companions or any(p.suffix.lower() == '.pptx' for p in references)):
        raise ValueError('Production references must be images only; PDF/PPTX joint learning uses reference_learning')
    paths = [pdf, *references, *map(Path, companions.values())]
    if request is not None:
        paths.append(Path(request))
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
    if pdf.suffix.lower() != '.pdf':
        raise ValueError('Missing-figure paper must be a PDF')
    for i, ref in enumerate(references, 1):
        if ref.suffix.lower() not in IMAGE_SUFFIXES | {'.pptx'}:
            raise ValueError('Reference must be PNG/JPEG/WebP or PPTX')
        if i in companions and (ref.suffix.lower() == '.pptx' or
                                Path(companions[i]).suffix.lower() != '.pptx'):
            raise ValueError('Native companion must be PPTX paired with an image reference')
    # Validate readability before creating the run; inspect original aspect, never resize.
    from PIL import Image
    with pymupdf.open(pdf) as doc:
        paper_text = '\n\n'.join(f'PAGE {i+1}\n{page.get_text()}'
                                  for i, page in enumerate(doc))
    image_sizes = {}
    for i, ref in enumerate(references, 1):
        if ref.suffix.lower() == '.pptx':
            Presentation(ref)
        else:
            with Image.open(ref) as img:
                image_sizes[i] = list(img.size)
                img.verify()
    for path in companions.values():
        Presentation(path)
    root.mkdir(parents=True, exist_ok=False)
    inputs = root / 'inputs'
    inputs.mkdir()
    shutil.copy2(pdf, inputs / 'paper-without-main-figure.pdf')
    entries = []
    for i, (ref, identity) in enumerate(zip(references, identities), 1):
        stem = 'reference' if len(references) == 1 else f'reference-{i:02d}'
        prefix = 'style-reference' if len(references) == 1 else f'style-reference-{i:02d}'
        asset_name = prefix + ref.suffix.lower()
        shutil.copy2(ref, inputs / asset_name)
        native_name = asset_name if ref.suffix.lower() == '.pptx' else None
        if i in companions:
            native_name = prefix + '-native.pptx'
            shutil.copy2(companions[i], inputs / native_name)
        entries.append({'index': i, 'paper_id': identity, 'asset': asset_name,
                        'kind': 'pptx' if ref.suffix.lower() == '.pptx' else 'image',
                        'pptx': native_name, 'image_pixels': image_sizes.get(i),
                        'visual_evidence': stem + '-render' if native_name else asset_name,
                        'pairing': 'caller-declared; model checks image/native correspondence'
                                   if i in companions else None})
    (inputs / 'request.txt').write_text(Path(request).read_text()
                                       if request is not None else DEFAULT_REQUEST)
    contract = {'aspect_ratio': [16, 9], 'slide_count': 1,
                'editable_primary': 'pptx', 'exports': ['pdf', 'png'],
                'global_presentation_title': False}
    write(root / 'inputs.json', {'schema': 'model-led-drawing.v2',
          'mode': mode, 'reference_input_policy': 'images_only' if mode == 'production' else 'legacy_native_inspection',
          'target_paper_id': target_id, 'reference_paper_id': identities[0],
          'style_references': entries, 'figure_contract': contract,
          'different_paper_identity': 'caller declaration when supplied; otherwise model inspection required',
          'hashes': {p.name: digest(p) for p in inputs.iterdir()},
          'outer_loop': outer_loop, 'model_executor': 'active session', 'quality': 'unassessed'})
    for entry in entries:
        if not entry['pptx']:
            continue
        native = inputs / entry['pptx']
        stem = 'reference' if len(entries) == 1 else f"reference-{entry['index']:02d}"
        write(inputs / f'{stem}-native-objects.json', inventory(native))
        write(inputs / f'{stem}-object-evidence.json', object_evidence(native))
        render(native, inputs / f'{stem}-render')
    (inputs / 'paper-text.txt').write_text(paper_text)
    workflow = guide_base / 'MODEL_LED_DRAWING.zh-CN.md'
    shutil.copy2(workflow, root / 'model-instructions.md')
    shutil.copy2(workflow.with_name('OBJECT_COMPOSITION.zh-CN.md'),
                 root / 'OBJECT_COMPOSITION.zh-CN.md')
    shutil.copy2(workflow.with_name('STYLE_LEARNING.zh-CN.md'),
                 root / 'STYLE_LEARNING.zh-CN.md')
    shutil.copy2(workflow.with_name('REFERENCE_JOINT_LEARNING.zh-CN.md'),
                 root / 'REFERENCE_JOINT_LEARNING.zh-CN.md')
    write(root / 'input-evidence.json', {'hashes': {
        str(p.relative_to(root)): digest(p)
        for p in [*inputs.rglob('*'), root / 'model-instructions.md',
                  root / 'OBJECT_COMPOSITION.zh-CN.md',
                  root / 'STYLE_LEARNING.zh-CN.md',
                  root / 'REFERENCE_JOINT_LEARNING.zh-CN.md'] if p.is_file()},
        'scope': 'byte consistency, not semantic provenance or OS isolation'})
    (root / 'iterations').mkdir()
    (root / 'human-feedback').mkdir()
    return str(root)


def verify_model_evidence(folder):
    """Verify model-authored notes are unchanged; never assess their meaning."""
    launch = folder / 'launch.json'
    if not launch.exists():
        return {}
    hashes = json.loads(launch.read_text()).get('model_evidence_hashes', {})
    for name, sha in hashes.items():
        if digest(folder / name) != sha:
            raise ValueError('Model evidence changed: ' + name)
    return hashes


def execute(root, builder, argv, decision, model_evidence=None):
    """Execute model-authored unrestricted local code; retain every attempt."""
    manifest = verify(root)
    notes = {}
    for label, source in (model_evidence or {}).items():
        if (not re.fullmatch(r'[a-z][a-z0-9_-]{0,63}', label) or
                label in {'decision', 'review'}):
            raise ValueError('Model evidence needs a unique plain label')
        source = Path(source)
        if not source.read_text().strip():
            raise ValueError('Model evidence cannot be empty: ' + label)
        notes['model-' + label + '.md'] = source
    attempts = root / 'iterations'
    iteration = attempts / f'{len(list(attempts.iterdir())):03d}'
    iteration.mkdir()
    shutil.copy2(builder, iteration / Path(builder).name)
    shutil.copy2(decision, iteration / 'model-decision.md')
    for name, source in notes.items():
        shutil.copy2(source, iteration / name)
    command = [str(x).replace('{builder}', str((iteration / Path(builder).name).resolve()))
               .replace('{output}', str(iteration.resolve())) for x in argv]
    write(iteration / 'launch.json', {'argv': command, 'builder_sha256': digest(builder),
          'started_at': datetime.now(timezone.utc).isoformat(),
          'decision_sha256': digest(decision), 'author': 'active model session',
          'model_evidence_hashes': {name: digest(iteration / name) for name in notes},
          'dependencies': 'local runtime, not hermetically frozen'})
    started = time.monotonic()
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=240)
        write(iteration / 'execution.json', {'exit_code': result.returncode,
              'stdout': result.stdout, 'stderr': result.stderr,
              'seconds': time.monotonic() - started})
        if result.returncode:
            raise RuntimeError('Model builder failed; inspect retained execution.json')
        verify(root)
        pptx = iteration / 'candidate.pptx'
        geometry = (figure_contract(pptx, manifest['figure_contract'])
                    if 'figure_contract' in manifest else None)
        write(iteration / 'native-objects.json', inventory(pptx))
        write(iteration / 'object-evidence.json', object_evidence(pptx))
        hashes = render(pptx, iteration / 'render')
        write(iteration / 'candidate.json', {'pptx_sha256': digest(pptx),
              'render_hashes': hashes, 'figure_contract_check': geometry,
              'quality': 'awaiting_model_visual_review',
              'human_acceptance': 'pending'})
    except Exception as error:
        write(iteration / 'failure.json', {'error': str(error),
                                          'automatic_retry': False})
        raise
    return str(iteration)


def record_review(root, iteration, review):
    verify(root)
    if len(iteration) != 3 or not iteration.isdigit():
        raise ValueError('Iteration must be a three-digit local index')
    folder = root / 'iterations' / iteration
    verify_model_evidence(folder)
    candidate = json.loads((folder / 'candidate.json').read_text())
    if digest(folder / 'candidate.pptx') != candidate['pptx_sha256']:
        raise ValueError('Candidate changed after render')
    for name, sha in candidate['render_hashes'].items():
        if digest(folder / 'render' / name) != sha:
            raise ValueError('Rendered evidence changed')
    # Free-form judgment is deliberately not reduced to a fixed pass/fail schema.
    if not Path(review).read_text().strip():
        raise ValueError('Model visual review cannot be empty')
    destination = folder / 'model-review.md'
    with destination.open('x') as stream:
        stream.write(Path(review).read_text())
    write(folder / 'review-record.json', {'review_sha256': digest(destination),
          'candidate_sha256': candidate['pptx_sha256'], 'author': 'active model session',
          'status': 'model_review_recorded', 'human_acceptance': 'pending',
          'visual_inspection': 'model attestation in review, not mechanically verified'})
    return str(folder / 'model-review.md')


def deliver(root, iteration, out, caption):
    """Publish the model-selected reviewed candidate; no automatic ranking/acceptance."""
    manifest = verify(root)
    if len(iteration) != 3 or not iteration.isdigit():
        raise ValueError('Iteration must be a three-digit local index')
    folder = root / 'iterations' / iteration
    model_notes = verify_model_evidence(folder)
    candidate = json.loads((folder / 'candidate.json').read_text())
    record = json.loads((folder / 'review-record.json').read_text())
    if (digest(folder / 'candidate.pptx') != candidate['pptx_sha256'] or
        record['candidate_sha256'] != candidate['pptx_sha256'] or
        digest(folder / 'model-review.md') != record['review_sha256']):
        raise ValueError('Candidate or model review changed after recording')
    for name, sha in candidate['render_hashes'].items():
        if digest(folder / 'render' / name) != sha:
            raise ValueError('Rendered evidence changed')
    launch = folder / 'launch.json'
    if launch.exists() and digest(folder / 'model-decision.md') != json.loads(
            launch.read_text())['decision_sha256']:
        raise ValueError('Model design evidence changed after execution')
    geometry = figure_contract(folder / 'candidate.pptx',
                               manifest.get('figure_contract', {'aspect_ratio': [16, 9]}))
    pdf = folder / 'render' / 'candidate.pdf'
    png = folder / 'render' / 'slide-01.png'
    if not {'candidate.pdf', 'slide-01.png'} <= candidate['render_hashes'].keys():
        raise ValueError('Delivery requires recorded actual PDF and PNG exports')
    with pymupdf.open(pdf) as doc:
        if len(doc) != 1 or abs(doc[0].rect.width / doc[0].rect.height - 16/9) > 1e-4:
            raise ValueError('Actual PDF export must have one 16:9 page')
    from PIL import Image
    with Image.open(png) as img:
        if abs(img.width * 9 - img.height * 16) > 16:
            raise ValueError('Actual PNG export must be 16:9 within pixel rounding')
    caption = Path(caption)
    if not caption.is_file() or not caption.read_text().strip():
        raise ValueError('Model-authored external figure caption is required')
    out = Path(out)
    out.mkdir(parents=True, exist_ok=False)
    for source, destination in [(folder / 'candidate.pptx', 'main-figure.pptx'),
            (pdf, 'main-figure.pdf'), (png, 'main-figure.png'),
            (caption, 'figure-caption.md'), (folder / 'model-decision.md', 'design-evidence.md'),
            (folder / 'model-review.md', 'visual-review.md')]:
        shutil.copy2(source, out / destination)
    for name in model_notes:
        shutil.copy2(folder / name, out / name)
    proxy = folder / 'render' / 'slide-01-publication.png'
    if proxy.name in candidate['render_hashes']:
        shutil.copy2(proxy, out / 'publication-preview.png')
    else:
        # Older records have only the full-size PNG; preserve their audit boundary.
        with pymupdf.open(pdf) as doc:
            page = doc[0]
            page.get_pixmap(matrix=pymupdf.Matrix(1063/page.rect.width,
                                                  1063/page.rect.width)).save(out / 'publication-preview.png')
    write(out / 'delivery.json', {'schema': 'model-led-drawing-delivery.v2',
          'iteration': iteration, 'figure_contract_check': geometry,
          'hashes': {p.name: digest(p) for p in out.iterdir()},
          'selected_by': 'active model, explicit iteration', 'human_acceptance': 'pending',
          'quality': 'model review recorded, not mechanically established',
          'publication_proxy': '180mm width at 150dpi, inspect visually before handoff'})
    return str(out)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('prepare')
    for key in ('root', 'pdf'):
        p.add_argument('--' + key, type=Path, required=True)
    p.add_argument('--request', type=Path)
    p.add_argument('--reference', type=Path, action='append', required=True,
                   help='Repeat for n reference images; PPTX allowed only in explicit legacy mode')
    p.add_argument('--reference-native', action='append', default=[], metavar='INDEX=PPTX',
                   help='Native companion for explicit legacy inspection only')
    p.add_argument('--mode', choices=['production', 'legacy-reference-inspection'], default='production')
    p.add_argument('--reference-id', action='append')
    p.add_argument('--target-id')
    p.add_argument('--guide-directory', type=Path,
                   help='Explicit versioned general guides; never per-sample teacher material')
    p.add_argument('--outer-loop', choices=['human', 'model'], default='human')
    p = sub.add_parser('execute')
    for key in ('root', 'builder', 'decision'):
        p.add_argument('--' + key, type=Path, required=True)
    p.add_argument('--argv', required=True, help='JSON argv using {builder} and {output}')
    p.add_argument('--model-evidence', action='append', default=[], metavar='LABEL=MARKDOWN',
                   help='Retain free-form model understanding/representation notes, not a scene schema')
    p = sub.add_parser('review')
    p.add_argument('--root', type=Path, required=True)
    p.add_argument('--iteration', required=True)
    p.add_argument('--review', type=Path, required=True)
    p = sub.add_parser('deliver')
    for key in ('root', 'out', 'caption'):
        p.add_argument('--' + key, type=Path, required=True)
    p.add_argument('--iteration', required=True)
    args = parser.parse_args(argv)
    if args.command == 'prepare':
        companions = {}
        for item in args.reference_native:
            try:
                index, path = item.split('=', 1)
                index = int(index)
                if index in companions:
                    raise ValueError('Duplicate native companion index')
                companions[index] = Path(path)
            except ValueError as error:
                parser.error(str(error) + '; use INDEX=PPTX')
        result = prepare(args.root.resolve(), args.pdf, args.reference,
                         args.target_id, args.reference_id, args.request, companions, args.mode,
                         args.guide_directory, args.outer_loop)
    elif args.command == 'execute':
        notes = {}
        for item in args.model_evidence:
            label, separator, source = item.partition('=')
            if not separator or not source or label in notes:
                parser.error('Use a unique LABEL=MARKDOWN for model evidence')
            notes[label] = Path(source)
        result = execute(args.root.resolve(), args.builder, json.loads(args.argv), args.decision, notes)
    elif args.command == 'review':
        result = record_review(args.root.resolve(), args.iteration, args.review)
    else:
        result = deliver(args.root.resolve(), args.iteration, args.out.resolve(), args.caption)
    print(result)


if __name__ == '__main__':
    main()
