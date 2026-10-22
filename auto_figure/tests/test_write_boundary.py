"""The builder's ordinary writes cannot affect sibling files; no system changes."""
import subprocess
from pathlib import Path
from auto_figure.figurecrafter_runner import ROOT, python_path


def test_landlock_blocks_ordinary_writes_outside_generation(tmp_path):
    allowed = tmp_path / 'allowed'
    allowed.mkdir()
    protected = tmp_path / 'protected.txt'
    protected.write_text('unchanged')
    script = '''import sys
from pathlib import Path
sys.path.insert(0,sys.argv[1])
from auto_figure.write_boundary import restrict_writes
r=restrict_writes([sys.argv[2]])
Path(sys.argv[2],'ok.txt').write_text('allowed')
try:
 Path(sys.argv[3]).write_text('bad')
except PermissionError:
 print('blocked',r['landlock_abi'])
else:
 raise SystemExit('ordinary protected write was allowed')
'''
    result = subprocess.run([python_path(), '-c', script, str(ROOT.parent), str(allowed), str(protected)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert 'blocked' in result.stdout
    assert protected.read_text() == 'unchanged'
    assert (allowed/'ok.txt').read_text() == 'allowed'


def test_real_libreoffice_render_with_project_write_boundary(tmp_path):
    from auto_figure.figurecrafter_runner import runtime_environment, REPO
    code = '''import sys
from pathlib import Path
sys.path.insert(0,sys.argv[1]);sys.path.insert(0,sys.argv[2])
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from auto_figure.write_boundary import restrict_writes
from harness.model_led_drawing import render
root=Path(sys.argv[3]);repo=Path(sys.argv[1])
restrict_writes([root,repo/'.cache/figurecrafter',repo/'.cache/fc-ipc',repo/'assets/tmp/auto_figure'])
p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9)
s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(1),Inches(4),Inches(3)).text='Editable fixture'
p.save(root/'probe.pptx');render(root/'probe.pptx',root/'render')
assert (root/'render/probe.pdf').is_file()
print('actual confined render passed')
'''
    result=subprocess.run([python_path(),'-c',code,str(REPO),str(ROOT/'figurecrafter'),str(tmp_path)],capture_output=True,text=True,env=runtime_environment(),timeout=120)
    assert result.returncode==0, result.stderr
    assert 'actual confined render passed' in result.stdout
